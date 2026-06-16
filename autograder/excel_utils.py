"""Excel utilities – read student roster (.xls/.xlsx) and export graded results."""

from __future__ import annotations

from pathlib import Path
from typing import Any

import openpyxl
import xlrd
import xlwt


def read_student_roster(path: str | Path) -> list[dict[str, Any]]:
    """Read the student roster spreadsheet.

    Supports both .xls (xlrd) and .xlsx (openpyxl).
    Returns a list of dicts with keys from the header row.
    """
    path = Path(path)
    if path.suffix == ".xls":
        return _read_xls(path)
    return _read_xlsx(path)


def _read_xls(path: Path) -> list[dict[str, Any]]:
    wb = xlrd.open_workbook(str(path))
    ws = wb.sheet_by_index(0)
    headers = [str(ws.cell_value(0, c)).strip() for c in range(ws.ncols)]
    rows: list[dict[str, Any]] = []
    for r in range(1, ws.nrows):
        row = {}
        for c, h in enumerate(headers):
            val = ws.cell_value(r, c)
            if isinstance(val, float) and val == int(val):
                val = int(val)
            row[h] = val
        rows.append(row)
    return rows


def _read_xlsx(path: Path) -> list[dict[str, Any]]:
    wb = openpyxl.load_workbook(str(path), read_only=True)
    ws = wb.active
    data = list(ws.iter_rows(values_only=True))
    if not data:
        return []
    headers = [str(h).strip() for h in data[0]]
    rows: list[dict[str, Any]] = []
    for row_vals in data[1:]:
        row = {}
        for h, v in zip(headers, row_vals):
            row[h] = v
        rows.append(row)
    wb.close()
    return rows


HEADER_SCORE = "成绩（录入项）"
HEADER_COMMENT = "评语（录入项）"
HEADER_STUDENT_ID = "学号"


def _ensure_output_headers(headers: list[str]) -> list[str]:
    out = list(headers)
    for h in (HEADER_SCORE, HEADER_COMMENT):
        if h not in out:
            out.append(h)
    return out


def export_graded_xlsx(
    roster_path: str | Path,
    results: dict[str, tuple[int, str]],
    out_path: str | Path,
) -> None:
    """Copy the roster and fill in scores/comments.

    *results* maps student_id -> (total_score, comment_text).
    """
    roster_path = Path(roster_path)
    out_path = Path(out_path)

    if roster_path.suffix == ".xls":
        _export_from_xls(roster_path, results, out_path)
    else:
        _export_from_xlsx(roster_path, results, out_path)


def _export_from_xls(src: Path, results: dict, dst: Path) -> None:
    """Read .xls with xlrd, write .xls (xlwt) 或 .xlsx (openpyxl) 按 dst 后缀。"""
    rows = _read_xls(src)
    if not rows:
        if dst.suffix.lower() == ".xls":
            wb = xlwt.Workbook()
            wb.add_sheet("Sheet1").write(0, 0, HEADER_STUDENT_ID)
            wb.save(str(dst))
        else:
            openpyxl.Workbook().save(str(dst))
        return

    headers = _ensure_output_headers(list(rows[0].keys()))
    for row in rows:
        sid = str(row.get(HEADER_STUDENT_ID, ""))
        if sid in results:
            row[HEADER_SCORE], row[HEADER_COMMENT] = results[sid]

    if dst.suffix.lower() == ".xls":
        wb = xlwt.Workbook()
        ws = wb.add_sheet("Sheet1")
        for c, h in enumerate(headers):
            ws.write(0, c, h)
        for r_idx, row in enumerate(rows, 1):
            for c, h in enumerate(headers):
                val = row.get(h, "")
                ws.write(r_idx, c, val)
        wb.save(str(dst))
        return

    wb = openpyxl.Workbook()
    ws = wb.active
    for c, h in enumerate(headers, 1):
        ws.cell(1, c, h)
    for r_idx, row in enumerate(rows, 2):
        for c, h in enumerate(headers, 1):
            ws.cell(r_idx, c, row.get(h, ""))
    wb.save(str(dst))


def _export_from_xlsx(src: Path, results: dict, dst: Path) -> None:
    wb = openpyxl.load_workbook(str(src))
    ws = wb.active
    headers = [str(c.value).strip() for c in ws[1]]

    sid_col = headers.index(HEADER_STUDENT_ID) + 1 if HEADER_STUDENT_ID in headers else None
    if HEADER_SCORE not in headers:
        headers.append(HEADER_SCORE)
        ws.cell(1, len(headers), HEADER_SCORE)
    if HEADER_COMMENT not in headers:
        headers.append(HEADER_COMMENT)
        ws.cell(1, len(headers), HEADER_COMMENT)
    score_col = headers.index(HEADER_SCORE) + 1
    comment_col = headers.index(HEADER_COMMENT) + 1

    if sid_col is None:
        wb.save(str(dst))
        return

    for row in ws.iter_rows(min_row=2):
        sid = str(row[sid_col - 1].value or "")
        if sid in results:
            score, comment = results[sid]
            if score_col:
                ws.cell(row[0].row, score_col, score)
            if comment_col:
                ws.cell(row[0].row, comment_col, comment)
    wb.save(str(dst))


def get_roster_headers(in_path: str | Path) -> list[str]:
    """读取 in.xls 的表头，保证导出列与输入完全一致。"""
    in_path = Path(in_path)
    if not in_path.exists():
        return [HEADER_STUDENT_ID, "姓名", HEADER_SCORE, HEADER_COMMENT]
    rows = read_student_roster(in_path)
    if not rows:
        return [HEADER_STUDENT_ID, "姓名", HEADER_SCORE, HEADER_COMMENT]
    return list(rows[0].keys())


def write_roster_to_xls(
    roster: list[dict[str, Any]],
    headers: list[str],
    out_path: str | Path,
) -> None:
    """将「评分统计」表格按 in.xls 的列顺序写入 .xls，供下载使用。"""
    out_path = Path(out_path)
    headers = _ensure_output_headers(headers)
    wb = xlwt.Workbook()
    ws = wb.add_sheet("Sheet1")
    for c, h in enumerate(headers):
        ws.write(0, c, h)
    for r_idx, row in enumerate(roster):
        for c, h in enumerate(headers):
            val = row.get(h, "")
            ws.write(r_idx + 1, c, val)
    wb.save(str(out_path))
